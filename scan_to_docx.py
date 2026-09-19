"""Turn a scanned, image-only PDF into an editable Word document.

Prose, headings, bullets, inline bold, ruled tables, multi-column pages and
figures. Every table is also written to a parallel spreadsheet. Every run
writes a quality report naming what it repaired and what a human must check.

    python scan_to_docx.py input.pdf
    python scan_to_docx.py input.pdf --exclude-pages 1,36-37 --jobs 8
    python scan_to_docx.py input.pdf --config job.json

Requires Tesseract on PATH or at the default Windows install location.
"""

import argparse
import datetime as dt
import json
import os
import re
import shutil
import tempfile
import time
import traceback
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

import numpy as np
import pytesseract
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Emu, Inches, Pt

import pagelab as P
import pageworker as W
import tablebook
from qualitycheck import Checker, id_gaps, orphan_ids, write_report

TESSERACT_FALLBACKS = [
    r"C:\Users\{user}\AppData\Local\Programs\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
]

PAGE_NO = re.compile(r"^\s*page\s*\|?\s*\d*\s*(of\s*\d+)?\s*[|.]?\s*$", re.I)
NUMBERED = re.compile(r"^\d+\.?\s+\S")
SUBNUMBERED = re.compile(r"^\d+\.\d+")
WORD = re.compile(r"[A-Za-z]{3,}")
ALNUM = re.compile(r"[A-Za-z0-9]")
CLEAN = re.compile(r"[A-Za-z0-9 .,;:()/&%'\"\-]")
BULLET = re.compile(r"^[@®©e\u2022\u25cf\u25cb*]\s+")
REQ_ID = re.compile(r"^([A-Z])\s*[-\u2014~]?\s*([O0-9]{1,3})$", re.I)

BOLD_LINE_AT = 1.15  # line stroke weight relative to the document's normal
BOLD_WORD_AT = 1.30  # a single word inside an unemphasised line
ITALIC_LINE_AT = 0.90  # extra rightward shear; deliberately high, see MANUAL
LONG_ENOUGH = 4  # shorter words are too small to judge on their own
USABLE_WIDTH = Inches(6.3)


# ---------------------------------------------------------------- emphasis


def weigh_lines(lines):
    """Normalise each word's stroke width by its own line's height.

    Dividing by the word's own height, the obvious move, is wrong: a word with
    no ascenders is shorter at the same weight, so 'or' and 'a' come out bold.
    """
    for line in lines:
        solid = [w for w in line["words"] if w["style"] and len(w["t"]) >= LONG_ENOUGH]
        pool = solid or [w for w in line["words"] if w["style"]]
        if not pool:
            continue
        height = float(np.median([w["h"] for w in pool])) or 1.0
        for w in line["words"]:
            if w["style"]:
                w["style"]["weight"] = w["style"]["stroke"] / height


def apply_emphasis(lines, want_italic=False):
    """Mark words heavier or more slanted than the document's normal text."""
    weigh_lines(lines)
    judged = [
        w
        for l in lines
        for w in l["words"]
        if w["style"] and len(w["t"]) >= LONG_ENOUGH
    ]
    if len(judged) < 20:
        return 0, 0

    weight = float(np.median([w["style"]["weight"] for w in judged]))
    slant = float(np.median([w["style"]["slant"] for w in judged]))
    if weight <= 0:
        return 0, 0

    for line in lines:
        ws = line["words"]
        solid = [w for w in ws if w["style"] and len(w["t"]) >= LONG_ENOUGH]
        if not solid:
            continue
        # The line decides first. Whole-line emphasis is what scans actually
        # carry: headings, a bold clause, a table header. A single word is only
        # called bold on its own when it is markedly heavier than that.
        line_ratio = np.median([w["style"]["weight"] for w in solid]) / weight
        line_slant = np.median([w["style"]["slant"] for w in solid]) - slant
        line_bold = line_ratio >= BOLD_LINE_AT
        line_italic = want_italic and line_slant >= ITALIC_LINE_AT

        for w in ws:
            if line_bold:
                w["bold"] = True
            elif w["style"] and len(w["t"]) >= LONG_ENOUGH:
                w["bold"] = w["style"]["weight"] / weight >= BOLD_WORD_AT
            w["italic"] = line_italic

        # A short word between two emphasised words belongs with them.
        for i, w in enumerate(ws):
            if w["style"] and len(w["t"]) >= LONG_ENOUGH:
                continue
            left = next(
                (x for x in reversed(ws[:i]) if len(x["t"]) >= LONG_ENOUGH), None
            )
            right = next((x for x in ws[i + 1 :] if len(x["t"]) >= LONG_ENOUGH), None)
            near = [x["bold"] for x in (left, right) if x is not None]
            w["bold"] = bool(near) and all(near)

    bold = sum(1 for l in lines for w in l["words"] if w["bold"])
    italic = sum(1 for l in lines for w in l["words"] if w["italic"])
    return bold, italic


# ------------------------------------------------------------------- lines


def lines_from(words):
    """Group words into lines, keeping the words so emphasis survives."""
    grouped = {}
    for w in words:
        grouped.setdefault(w["line"], []).append(w)
    out = []
    for g in grouped.values():
        g.sort(key=lambda w: w["x"])
        out.append(
            {
                "words": g,
                "text": " ".join(w["t"] for w in g).strip(" |"),
                "x": min(w["x"] for w in g),
                "y": min(w["y"] for w in g),
                "h": float(np.median([w["h"] for w in g])),
                "conf": float(np.mean([w["c"] for w in g])),
            }
        )
    return sorted(out, key=lambda l: l["y"])


def is_english(ln):
    """Reject other scripts, which English-only OCR returns as noise."""
    txt = ln["text"]
    if not txt or not ALNUM.search(txt):
        return False
    if len(CLEAN.findall(txt)) / len(txt) < 0.75:
        return False
    if len(txt) > 12 and not WORD.search(txt):
        return False
    return ln["conf"] >= (30 if len(txt) <= 6 else 55)


# ------------------------------------------------------------------- tables


def cell_lines(words, x0, x1, y0, y1):
    inside = [
        w
        for w in words
        if x0 <= w["x"] + w["w"] / 2 <= x1 and y0 <= w["y"] + w["h"] / 2 <= y1
    ]
    return [l for l in lines_from(inside) if is_english(l)]


def tidy_id(txt):
    """Repair reference IDs the OCR reads as M-O1, M- 03, D-02 and so on."""
    m = REQ_ID.match(txt.strip())
    if not m:
        return txt
    return f"{m.group(1).upper()}-{m.group(2).upper().replace('O', '0'):0>2}"


def grid(table, words):
    """A table's cells, each a list of lines, so emphasis survives."""
    data = []
    for y0, y1 in zip(table["rows"], table["rows"][1:]):
        row = [
            cell_lines(words, x0, x1, y0, y1)
            for x0, x1 in zip(table["cols"], table["cols"][1:])
        ]
        if any(cell for cell in row):
            data.append(row)
    return data


def cell_string(cell):
    return " ".join(l["text"] for l in cell).strip()


def set_widths(tbl, cols):
    """Give the Word table the column proportions the scan actually had."""
    spans = [b - a for a, b in zip(cols, cols[1:])]
    total = float(sum(spans)) or 1.0
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, span in enumerate(spans):
        width = Emu(int(USABLE_WIDTH * (span / total)))
        for row in tbl.rows:
            row.cells[j].width = width


# --------------------------------------------------------------- page order


def trustworthy(known, span):
    """Footer readings worth acting on.

    Scans go wrong locally: pages get transposed or slipped by a place or two,
    never randomly permuted. A reading far from its own position is a misread
    footer, and a number claimed by two pages is evidence of nothing.
    """
    near = [(i, n) for i, n in known if abs(n - (i + 1)) <= span]
    seen = {}
    for _, n in near:
        seen[n] = seen.get(n, 0) + 1
    return {i: n for i, n in near if seen[n] == 1}


def reorder(items, numbers, span=3):
    """Sort pages by printed number, keeping unnumbered pages where they are."""
    known = [(i, n) for i, n in enumerate(numbers) if n is not None]
    trusted = trustworthy(known, span)
    rejected = [(i + 1, n) for i, n in known if trusted.get(i) != n]
    if not trusted or all(n == i + 1 for i, n in trusted.items()):
        return items, [], rejected

    keys, last = {}, 0.0
    for i in range(len(items)):
        if i in trusted:
            last = float(trusted[i])
        else:
            last += 0.001  # hold position behind the last known page
        keys[i] = last

    moved = [(i + 1, n) for i, n in trusted.items() if n != i + 1]
    order = sorted(range(len(items)), key=lambda i: keys[i])
    return [items[i] for i in order], moved, rejected


# ------------------------------------------------------------------ writing


def add_runs(paragraph, lines, checker, where, strip_bullet=False, mark=True):
    """Write lines into a paragraph, one run per stretch of common styling.

    Words the OCR was unsure about are highlighted, so the checking happens in
    the document itself rather than against a list of line numbers.
    """
    words = [w for l in lines for w in l["words"]]
    run, state = None, None
    for i, w in enumerate(words):
        text = checker.clean(w["t"], where)
        if strip_bullet and i == 0:
            text = BULLET.sub("", text, count=1)
        style = (w["bold"], w["italic"], bool(w.get("low")) and mark)
        if style != state or run is None:
            run = paragraph.add_run(("" if run is None else " ") + text)
            run.bold, run.italic = style[0], style[1]
            if style[2]:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            state = style
        else:
            run.text += " " + text
    return paragraph


def fill_cell(cell, lines, checker, page, override=None, mark=True):
    if override is not None:
        cell.text = override
        return
    if not lines:
        return
    cell.text = ""
    add_runs(cell.paragraphs[0], lines, checker, f"p{page + 1} table", mark=mark)


# ------------------------------------------------------------------ assembly


def assemble(doc, pages, checker, args, body_h, figure_dir):
    """Write every page into the document in reading order."""
    state = {
        "paragraphs": 0,
        "tables": [],
        "sheets": [],
        "figures": 0,
        "columns": [],
        "gaps": [],
        "orphans": [],
        "empty_rows": 0,
        "low_confidence": [],
        "schema": [],
        "highlighted": 0,
    }
    carry, carry_cols = None, None

    for page in pages:
        tables, words = page["tables"], page["words"]
        spans = [(t["top"], t["bot"]) for t in tables]
        prose = [
            w
            for w in words
            if not any(a - 8 <= w["y"] + w["h"] / 2 <= b + 8 for a, b in spans)
        ]
        text_lines = [
            l
            for l in lines_from(prose)
            if not PAGE_NO.match(l["text"]) and is_english(l)
        ]

        columns = [(0, page["width"])]
        if not page.get("single_column"):
            columns = P.column_bounds(
                [(w["x"], w["y"], w["w"], w["h"]) for w in prose], page["width"]
            )
        if len(columns) > 1:
            state["columns"].append(page["index"] + 1)

        def column_of(x):
            for n, (a, b) in enumerate(columns):
                if a <= x < b:
                    return n
            return 0

        elements = [("table", t["top"], 0, t) for t in tables]
        for line in text_lines:
            elements.append(("line", line["y"], column_of(line["x"]), line))
            if line["conf"] < W.LOW_CONFIDENCE:
                state["low_confidence"].append(
                    (page["index"] + 1, line["conf"], line["text"])
                )
        for fig in page["figures"]:
            elements.append(("figure", fig["y"], column_of(fig["x"]), fig))
        elements.sort(key=lambda e: (e[2], e[1]))

        carry, carry_cols = write_page(
            doc,
            elements,
            checker,
            state,
            body_h,
            carry,
            carry_cols,
            figure_dir,
            page,
            args,
        )
    return state


def write_page(
    doc, elements, checker, state, body_h, carry, carry_cols, figure_dir, page, args
):
    pending, style, prev_y = [], {"bullet": False}, None
    where = f"p{page['index'] + 1}"

    def flush():
        if not pending:
            return
        p = doc.add_paragraph(style="List Bullet" if style["bullet"] else None)
        add_runs(
            p,
            pending,
            checker,
            where,
            strip_bullet=style["bullet"],
            mark=not args.no_highlight,
        )
        if not style["bullet"]:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        state["paragraphs"] += 1
        state["highlighted"] += sum(
            1 for l in pending for w in l["words"] if w.get("low")
        )
        pending.clear()
        style["bullet"] = False

    for kind, _, _, payload in elements:
        if kind == "line":
            line = payload
            txt = line["text"]
            if line["h"] > body_h * 1.25 or SUBNUMBERED.match(txt):
                flush()
                doc.add_heading(
                    checker.clean(txt, where), level=1 if NUMBERED.match(txt) else 2
                )
                state["paragraphs"] += 1
                prev_y = None
                continue
            gap = line["y"] - prev_y if prev_y is not None else None
            marked = bool(BULLET.match(txt))
            # Inside a list the glyph is often lost to the scan, so a finished
            # sentence plus a fresh capital starts the next item.
            runs_on = (
                style["bullet"]
                and pending
                and pending[-1]["text"].rstrip().endswith(".")
                and txt[:1].isupper()
            )
            if gap is None or gap > body_h * 2.4 or marked or runs_on:
                was = style["bullet"]
                flush()
                style["bullet"] = marked or (runs_on and was)
            pending.append(line)
            prev_y = line["y"]
            continue

        flush()
        prev_y = None

        if kind == "figure":
            path = figure_dir / f"fig{state['figures']:03d}.png"
            path.write_bytes(payload["png"])
            doc.add_picture(str(path), width=Inches(min(6.3, payload["width"] / 150)))
            state["figures"] += 1
            continue

        data = grid(payload, page["words"])
        if not data:
            continue
        strings = [[cell_string(c) for c in row] for row in data]
        for row in strings:
            row[0] = tidy_id(row[0])

        if carry is not None:
            fit, why = continues(carry, carry_cols, data, strings, payload)
            if fit:
                if strings[0] == [c.text for c in carry.rows[0].cells]:
                    data, strings = data[1:], strings[1:]
                for row, text in zip(data, strings):
                    cells = carry.add_row().cells
                    for j, (cell, content) in enumerate(zip(cells, row)):
                        fill_cell(
                            cell,
                            content,
                            checker,
                            page["index"],
                            override=text[j] if j == 0 else None,
                            mark=not args.no_highlight,
                        )
                    state["sheets"][-1][0].append(text)
                continue
            if why:
                state["schema"].append((page["index"] + 1, why))

        tbl = doc.add_table(rows=len(data), cols=len(data[0]))
        tbl.style = "Table Grid"
        for i, row in enumerate(data):
            for j, content in enumerate(row):
                fill_cell(
                    tbl.cell(i, j),
                    content,
                    checker,
                    page["index"],
                    override=strings[i][j] if j == 0 else None,
                    mark=not args.no_highlight,
                )
        for cell in tbl.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
        set_widths(tbl, payload["cols"])
        state["sheets"].append(([list(r) for r in strings], page["index"] + 1))
        carry, carry_cols = tbl, payload["cols"]

    flush()
    # Only a table running to the foot of the page can continue overleaf.
    if not page["tables"] or page["tables"][-1]["bot"] < page["height"] * 0.85:
        carry, carry_cols = None, None
    return carry, carry_cols


def continues(carry, carry_cols, data, strings, table):
    """Whether this table is the previous one carrying on, and why not.

    Column count alone is a weak test: two unrelated five-column tables would
    be welded together. The column positions have to line up as well, and a
    header that half-agrees with the one above is decisive evidence they are
    different tables.
    """
    if len(data[0]) != len(carry.columns):
        return False, (
            f"continued with {len(data[0])} columns after "
            f"{len(carry.columns)}; started a new table instead"
        )

    header = [plain(c.text) for c in carry.rows[0].cells]
    first = [plain(s) for s in strings[0]]
    matches = sum(1 for a, b in zip(first, header) if a and a == b)
    # A repeated header that half agrees is two different tables. A first row
    # that agrees with nothing is simply data, and says nothing either way.
    if matches and matches * 2 < len(header):
        return False, (
            f"header {strings[0]} half agrees with "
            f"{[c.text for c in carry.rows[0].cells]}; "
            "started a new table instead"
        )

    if carry_cols and len(carry_cols) == len(table["cols"]):
        # Proportions, not pixels: a continued table keeps its column widths
        # even when the whole grid sits a little left or right on the page.
        drift = max(
            abs(a - b)
            for a, b in zip(proportions(carry_cols), proportions(table["cols"]))
        )
        if drift > 0.08:
            return False, (
                f"column widths differ by {drift:.0%} between "
                "pages; started a new table instead"
            )
    return True, None


def plain(text):
    """Text reduced to what OCR cannot easily get wrong."""
    return re.sub(r"[^a-z0-9]", "", text.lower())


def proportions(cols):
    span = float(cols[-1] - cols[0]) or 1.0
    return [(c - cols[0]) / span for c in cols]


# ------------------------------------------------------------------ plumbing


def locate_tesseract(explicit):
    if explicit:
        return explicit
    found = shutil.which("tesseract")
    if found:
        return found
    for pattern in TESSERACT_FALLBACKS:
        path = Path(pattern.format(user=Path.home().name))
        if path.exists():
            return str(path)
    raise SystemExit("Tesseract not found. Install it or pass --tesseract.")


def load_config(args):
    """Merge a JSON job file into the parsed arguments."""
    if not args.config:
        return args
    cfg = json.loads(Path(args.config).read_text(encoding="utf8"))
    for key, value in cfg.items():
        if not hasattr(args, key):
            raise SystemExit(f"Unknown setting in config: {key}")
        if value is not None:
            setattr(args, key, value)
    return args


def build_arguments():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument(
        "pdf",
        type=Path,
        nargs="+",
        help="one or more PDF files, and/or folders to scan for *.pdf",
    )

    io_grp = ap.add_argument_group("Input / output")
    io_grp.add_argument(
        "-o",
        "--out",
        type=Path,
        help="output .docx for a single file, or an output folder "
        "for several (default: a date-stamped folder; see --same-dir)",
    )
    io_grp.add_argument(
        "--same-dir",
        action="store_true",
        help="write output beside each source PDF instead of a " "date-stamped folder",
    )
    io_grp.add_argument("--config", type=Path, help="JSON job file of these settings")
    io_grp.add_argument(
        "--page-options",
        default={},
        help="per-page overrides, config file only, e.g. "
        '{"12": {"single_column": true, "no_tables": true}}',
    )
    io_grp.add_argument(
        "--no-xlsx",
        action="store_true",
        help="do not write the parallel spreadsheet of tables",
    )
    io_grp.add_argument(
        "--dictionary",
        type=Path,
        help="text file of extra words (one per line) the spell "
        "checker should accept, e.g. domain jargon or names",
    )

    excl_grp = ap.add_argument_group("Exclusions")
    excl_grp.add_argument(
        "--exclude-pages", default="", help="pages to leave out entirely, e.g. 1,36-37"
    )
    excl_grp.add_argument(
        "--exclude-template",
        action="append",
        default=[],
        type=Path,
        help="image of a region to remove wherever it appears",
    )
    excl_grp.add_argument(
        "--keep-furniture",
        action="store_true",
        help="keep the coloured letterhead and footer bands",
    )
    excl_grp.add_argument("--keep-stamp", action="store_true", help="keep ink stamps")

    layout_grp = ap.add_argument_group("Emphasis & layout")
    layout_grp.add_argument(
        "--no-reorder",
        action="store_true",
        help="trust the scan order instead of printed page numbers",
    )
    layout_grp.add_argument(
        "--no-emphasis", action="store_true", help="do not detect bold and italic"
    )
    layout_grp.add_argument(
        "--italic",
        action="store_true",
        help="also detect italics; slant is noisy, so this is off "
        "by default and will over-mark on some scans",
    )
    layout_grp.add_argument(
        "--no-figures", action="store_true", help="do not embed diagrams and images"
    )
    layout_grp.add_argument(
        "--no-autofix",
        action="store_true",
        help="report lost-space errors instead of repairing them",
    )
    layout_grp.add_argument(
        "--no-highlight",
        action="store_true",
        help="do not highlight low-confidence words in the document",
    )
    layout_grp.add_argument(
        "--no-second-pass",
        action="store_true",
        help="do not re-read low-confidence lines at higher zoom",
    )
    layout_grp.add_argument(
        "--no-deskew", action="store_true", help="do not straighten pages"
    )
    layout_grp.add_argument(
        "--no-despeckle", action="store_true", help="keep scanner dust"
    )

    perf_grp = ap.add_argument_group("Performance")
    perf_grp.add_argument(
        "--jobs", type=int, default=0, help="worker processes (default: half the cores)"
    )
    perf_grp.add_argument(
        "--no-cache", action="store_true", help="ignore cached page analysis"
    )
    perf_grp.add_argument(
        "--clear-cache", action="store_true", help="delete the cache before running"
    )
    perf_grp.add_argument("--dpi", type=int, default=300, help="render resolution")
    perf_grp.add_argument(
        "--scale", type=int, default=1, help="extra Lanczos upscale after rendering"
    )
    perf_grp.add_argument("--tesseract", help="path to tesseract.exe")
    return ap


def expand_inputs(paths):
    """Turn CLI paths (files and/or folders) into a sorted, deduplicated PDF list."""
    found = []
    for p in paths:
        if not p.exists():
            raise SystemExit(f"No such file or folder: {p}")
        if p.is_dir():
            found += sorted(p.glob("*.pdf"))
        else:
            found.append(p)
    seen, unique = set(), []
    for p in found:
        rp = p.resolve()
        if rp not in seen:
            seen.add(rp)
            unique.append(p)
    if not unique:
        raise SystemExit("No PDF files found.")
    return unique


def resolve_output_path(pdf, args, run_stamp):
    """Where one input's three output files should be written."""
    if args.out and len(args.pdf) == 1 and args.out.suffix.lower() == ".docx":
        return args.out
    if args.out:
        folder = args.out
    elif args.same_dir:
        folder = pdf.parent
    else:
        folder = pdf.parent / f"OCR_{run_stamp}"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / (pdf.stem + "_OCR.docx")


def convert_one(pdf, out_path, args, tess):
    """Convert one PDF. Returns a summary dict for both single and batch runs."""
    total = P.page_count(pdf)
    dropped = P.parse_pages(args.exclude_pages, total)
    page_options = args.page_options if isinstance(args.page_options, dict) else {}
    dropped |= {int(n) - 1 for n, o in page_options.items() if o.get("exclude")}

    options = {
        "dpi": args.dpi,
        "scale": args.scale,
        "tesseract": tess,
        "templates": [str(p) for p in args.exclude_template],
        "keep_furniture": args.keep_furniture,
        "keep_stamp": args.keep_stamp,
        "no_emphasis": args.no_emphasis,
        "no_figures": args.no_figures,
        "no_second_pass": args.no_second_pass,
        "no_deskew": args.no_deskew,
        "no_despeckle": args.no_despeckle,
        "page_options": page_options,
        "use_cache": not args.no_cache,
    }

    wanted = [i for i in range(total) if i not in dropped]
    workers = args.jobs or max(1, (os.cpu_count() or 2) // 2)
    print(
        f"{total} pages at {args.dpi} dpi, {len(wanted)} to convert, "
        f"{workers} worker(s)"
        + (f", excluding {sorted(n + 1 for n in dropped)}" if dropped else "")
    )

    started = time.time()
    jobs = [(str(pdf), i, options) for i in wanted]
    if workers > 1:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            analysed = list(pool.map(W.run_job, jobs, chunksize=1))
    else:
        analysed = [W.run_job(j) for j in jobs]
    analysed.sort(key=lambda p: p["index"])
    print(f"pages analysed in {time.time() - started:.0f}s")

    moved, rejected = [], []
    if not args.no_reorder:
        analysed, moved, rejected = reorder(analysed, [p["printed"] for p in analysed])
        if moved:
            print(
                "scan out of order; reading by printed number "
                "(scan page, printed page):",
                moved,
            )
        if rejected:
            print("footer readings ignored as inconsistent:", rejected)

    bold = italic = 0
    if not args.no_emphasis:
        for page in analysed:
            b, i = apply_emphasis(lines_from(page["words"]), args.italic)
            bold += b
            italic += i

    extra_words = []
    if args.dictionary:
        extra_words = [
            w.strip()
            for w in args.dictionary.read_text(encoding="utf-8").splitlines()
            if w.strip()
        ]
    checker = Checker(autofix=not args.no_autofix, extra_words=extra_words)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    body_h = float(np.median([w["h"] for p in analysed for w in p["words"]]) or 20)

    with tempfile.TemporaryDirectory() as tmp:
        run = assemble(doc, analysed, checker, args, body_h, Path(tmp))
        audit(doc, run)
        doc.save(out_path)

    sheets = 0
    book_path = out_path.with_name(out_path.stem + "_tables.xlsx")
    if not args.no_xlsx and run["sheets"]:
        sheets = tablebook.write_tables(book_path, run["sheets"], pdf.name)

    report_path = out_path.with_suffix(".quality.md")
    write_report(
        report_path,
        {
            "source": pdf,
            "output": out_path,
            "finished": dt.datetime.now(),
            "pages": len(analysed),
            "paragraphs": run["paragraphs"],
            "tables": run["tables"],
            "figures": run["figures"],
            "bold": bold,
            "italic": italic,
            "columns": run["columns"],
            "moved_pages": moved,
            "rejected_pages": rejected,
            "excluded": sorted(n + 1 for n in dropped),
            "templates": 0,
            "gaps": run["gaps"],
            "orphans": run["orphans"],
            "empty_rows": run["empty_rows"],
            "low_confidence": run["low_confidence"],
            "fixes": checker.fixes,
            "suspect": checker.suspect,
            "schema": run["schema"],
            "highlighted": run["highlighted"],
            "rescued": sum(p["rescued"] for p in analysed),
            "skewed": [
                (p["index"] + 1, round(p["skew"], 2))
                for p in analysed
                if abs(p["skew"]) >= 0.25
            ],
            "workbook": book_path.name if sheets else None,
            "elapsed": time.time() - started,
        },
    )

    elapsed = time.time() - started
    print(f"saved {out_path}")
    print(
        f"{run['paragraphs']} paragraphs, {len(run['tables'])} tables, "
        f"{run['figures']} figures, {bold} bold and {italic} italic runs, "
        f"{run['highlighted']} word(s) highlighted"
    )
    if sheets:
        print(f"tables also in {book_path} ({sheets} sheet(s))")
    print(f"quality report: {report_path}  ({elapsed:.0f}s)")

    return {
        "pdf": pdf,
        "out": out_path,
        "pages": len(analysed),
        "paragraphs": run["paragraphs"],
        "tables": len(run["tables"]),
        "figures": run["figures"],
        "highlighted": run["highlighted"],
        "elapsed": elapsed,
    }


def main():
    args = load_config(build_arguments().parse_args())
    inputs = expand_inputs(args.pdf)
    tess = locate_tesseract(args.tesseract)
    pytesseract.pytesseract.tesseract_cmd = tess
    if args.clear_cache:
        W.clear_cache()

    run_stamp = dt.date.today().isoformat()
    n = len(inputs)
    results, failed = [], []
    for i, pdf in enumerate(inputs, 1):
        out_path = resolve_output_path(pdf, args, run_stamp)
        print(f"[{i}/{n}] {pdf.name}")
        try:
            results.append(convert_one(pdf, out_path, args, tess))
        except Exception as exc:
            print(f"FAILED: {pdf.name}: {exc}")
            traceback.print_exc()
            failed.append((pdf, exc))
        print()

    if n > 1:
        ok = len(results)
        print(
            f"batch done: {ok}/{n} converted"
            + (f", {len(failed)} failed" if failed else "")
        )
        if results:
            print(
                f"  {sum(r['pages'] for r in results)} pages, "
                f"{sum(r['paragraphs'] for r in results)} paragraphs, "
                f"{sum(r['tables'] for r in results)} tables, "
                f"{sum(r['highlighted'] for r in results)} word(s) highlighted, "
                f"{sum(r['elapsed'] for r in results):.0f}s total"
            )
            folders = sorted({str(r["out"].parent) for r in results})
            print("  output in: " + ", ".join(folders))
        for pdf, exc in failed:
            print(f"  FAILED {pdf}: {exc}")
    print("Check figures, dates and clause numbers against the PDF before use.")
    if failed:
        raise SystemExit(1)


def audit(doc, state):
    """Re-measure the finished document: table shapes, gaps, empty rows."""
    state["tables"] = []
    for i, table in enumerate(doc.tables):
        rows = [[c.text for c in r.cells] for r in table.rows]
        header = " | ".join(c[:20] for c in rows[0])
        state["tables"].append((len(rows), len(rows[0]), header))
        state["empty_rows"] += sum(1 for r in rows[1:] if not any(c.strip() for c in r))
        for prefix, lo, hi, missing in id_gaps(rows):
            if missing:
                state["gaps"].append((i, prefix, lo, hi, missing))
        orphans = orphan_ids(rows)
        if orphans:
            state["orphans"].append((i, orphans))


if __name__ == "__main__":
    import sys

    if len(sys.argv) == 1:
        import interactive

        interactive.run()
    else:
        main()
