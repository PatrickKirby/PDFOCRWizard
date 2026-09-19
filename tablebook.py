"""Write every table the run found into a parallel spreadsheet.

A table in Word is for reading. The same table in a sheet is for working:
filtering a requirements matrix, pricing a bill of quantities, diffing two
revisions. Both come out of one pass, so neither has to be retyped.
"""

import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
UNSAFE = re.compile(r"[\\/*?:\[\]]")
NUMERIC = re.compile(r"^-?[\d,]+(\.\d+)?$")


def sheet_name(index, header, used):
    """A legible, unique, Excel-legal sheet name from the table's header."""
    base = UNSAFE.sub(" ", " ".join(header[:2])).strip() or f"Table {index + 1}"
    base = f"{index + 1}. {base}"[:31].strip()
    name, n = base, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def cast(value):
    """Numbers as numbers, so a spreadsheet can total a column of them."""
    text = value.strip()
    if NUMERIC.match(text):
        try:
            return (
                float(text.replace(",", ""))
                if "." in text
                else int(text.replace(",", ""))
            )
        except ValueError:
            return text
    return text


def write_tables(path, tables, source_name):
    """One sheet per table, plus a contents sheet pointing at each.

    `tables` is a list of (rows, page) where rows is a list of string lists.
    """
    book = Workbook()
    index = book.active
    index.title = "Contents"
    index.append(["Source", source_name])
    index.append([])
    index.append(["Table", "Page", "Rows", "Columns", "Header"])
    for cell in index[3]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    used = set()
    for i, (rows, page) in enumerate(tables):
        if not rows:
            continue
        name = sheet_name(i, rows[0], used)
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append([cast(c) for c in row])

        for cell in sheet[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"

        for col in range(1, len(rows[0]) + 1):
            longest = max(len(str(r[col - 1])) for r in rows)
            sheet.column_dimensions[get_column_letter(col)].width = min(
                60, max(12, longest * 0.95)
            )
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        index.append([name, page, len(rows), len(rows[0]), " | ".join(rows[0])[:120]])
        index.cell(index.max_row, 1).hyperlink = f"#'{name}'!A1"
        index.cell(index.max_row, 1).style = "Hyperlink"

    for col, width in zip("ABCDE", (34, 8, 8, 10, 80)):
        index.column_dimensions[col].width = width

    book.save(path)
    return len(book.sheetnames) - 1
